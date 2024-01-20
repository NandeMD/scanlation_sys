import os
from json import dumps
from orjson import loads

import discord
from discord import Attachment, Embed, File, Member, Role, EmbedField
from discord.ext import commands
from discord.ext.pages import Page, Paginator
from requests import get as rget
from xlsxwriter import Workbook
from docx import Document
from mp.misc import ext_from_filename, find_serie_by_id
from datetime import datetime
from io import BytesIO
from mp.scp_related import find_serie_and_chapter, update_chapter, update_serie, message_related_channel
import arrow

from mp.payment import calculate_pay, ChapterAndFees, extract_all_personnel_ids_chapters, match_personnel_ids_with_names_and_add_worksheet, prepare_sheet, split_all_work, generate_embeds_from_fields, split_a_page
from mp.config import get_api_ip, get_any


API = get_api_ip()
PR_FEE = get_any("PR_FEE")
QC_FEE = get_any("QC_FEE")
POOP_URL = get_any("POOP_URL")
DOLLAR_URL = get_any("DOLLAR_URL")


# noinspection PyUnresolvedReferences,DuplicatedCode
class MaasHesap(commands.Cog):
    def __init__(self, bot: commands.Bot):
        self.bot = bot
            
    @discord.slash_command(name="translate", description="Calculate the translation size, cost, and enter it into the bot's database.")
    async def translation(
        self, 
        ctx: discord.ApplicationContext, 
        tl_file: discord.Option(Attachment, name="tl_file", description="Your translation file should be in .txt or .docx format.", required=True),
        role: discord.Option(Role, name="serie_role", description="Serie role", required=True),
        chap_num: discord.Option(float, name="chapter_number", description="Chapter number. (like '7' or 7.5')", required=True),  # type: ignore
        pred: discord.Option(str, name="should_pred", description="PR needed?", choices=["Yes", "No"], required=True)
    ):

        serie, chapter = find_serie_and_chapter(role.id, chap_num)
        if not chapter:
            await ctx.respond("There is no such serie or chapter!", ephemeral=True)
            return
        
        attachment: Attachment = tl_file
        file_ext = ext_from_filename(attachment.filename)

        with rget(attachment.url) as file:
            data = file.content
            if file_ext == "txt":
                data = data.decode("utf-8")
            elif file_ext == "doc" or file_ext == "docx":
                doc = Document(BytesIO(file.content))
                data = "".join([p.text for p in doc.paragraphs])
            else:
                await ctx.respond("Unsupported file type!")
                return

        data = self.strip_data(data)
        time_str = datetime.now().strftime(r"%d/%m/%Y - %H.%M")

        serie_update_resp = update_serie(
            id=serie["id"],
            waiting_pr=chap_num,
            time_tl=time_str
        )

        if serie_update_resp.status_code != 200:
            await ctx.respond(f"An error occured while updating serie. Code: {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
            return

        if pred == "No":
            serie_update_resp2 = update_serie(
                id=serie["id"],
                pred=chap_num,
                time_pr=time_str
            )
            if serie_update_resp2.status_code != 200:
                await ctx.respond(f"An error occured while updating serie. Code: {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
                return

        chapter_update_resp = update_chapter(
            id=chapter["id"],
            tl_id=ctx.user.id,
            tl_name=ctx.user.name,
            tl_at=int(arrow.utcnow().timestamp()),
            tl_bytes=len(data),
            should_pred=1 if pred == "Yes" else 0
        )
        if chapter_update_resp.status_code != 200:
            await ctx.respond(f"An error occured while updating chapter. Code: {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
            return

        await ctx.respond(f"`{attachment.filename}` file is `{round(len(data) / 1000)}` kilobyte.\n"
                          f"Fee: {calculate_pay(len(data), 1 if pred == 'Yes' else 0)} $")
        await message_related_channel(self.bot, role, chap_num, "TL", ctx, pred, round(len(data) / 1000))
        
    @discord.slash_command(name="delete_tl", description="Delete the translation that you have entered to the bot.")
    async def deletetl(
        self, 
        ctx: discord.ApplicationContext, 
        role: discord.Option(Role, name="serie_role", description="Serie role.", required=True),
        chap_num: discord.Option(float, name="chapter_number", description="Chapter number. (like '7' or 7.5')", required=True)
    ):
        serie, chapter = find_serie_and_chapter(role.id, chap_num)

        if not serie:
            await ctx.respond("There is no such serie.")
            return
        if not chapter:
            await ctx.respond("There is no such TL.")
            return

        update_chapter_response = update_chapter(
            id=chapter["id"],
            tl_id=0,
            tl_name="",
            tl_at=0,
            tl_bytes=0,
            should_pred=0
        )
        if update_chapter_response.status_code != 200:
            await ctx.respond(f"An error occured while updating chapter. Code:  {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
            return

        await ctx.respond(f"`Chapter {role.name}` `{chap_num}` is deleted.")
    
    @discord.slash_command(name="ts", description="Calculate the typesetting fee and enter it into the bot's database.")
    async def calculatets(
        self, 
        ctx: discord.ApplicationContext,
        role: discord.Option(Role, name="serie_role", description="Serie role.", required=True),
        chap_num: discord.Option(float, name="chapter_num", description="Chapter number. (like '7' or 7.5')", required=True),
        pred: discord.Option(str, name="qc_needed", description="Is QC needed?", choices=["Yes", "No"], required=True)
    ):
        serie, chapter = find_serie_and_chapter(role.id, chap_num)
        if not chapter:
            await ctx.respond("There is no such serie or chapter!", ephemeral=True)
            return

        if not bool(chapter["tl_id"]):
            await ctx.respond("There is no translation found!", ephemeral=True)
            return

        time_str = datetime.now().strftime(r"%d/%m/%Y - %H.%M")

        update_serie_response = update_serie(
            id=serie["id"],
            completed=chap_num,
            time_ts=time_str
        )
        if update_serie_response.status_code != 200:
            await ctx.respond(f"An error occured. Code: {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
            return

        if pred == "No":
            update_serie_response2 = update_serie(
                id=serie["id"],
                last_qced=chap_num,
                time_qc=time_str
            )
            if update_serie_response2.status_code != 200:
                await ctx.respond(f"An error occured. Code: {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
                return
        if not len(str(serie["clnr"])) > 1 or not serie["clnr"]:
            update_serie_response3 = update_serie(
                id=serie["id"],
                cleaned=chap_num,
                time_cln=time_str
            )

        update_chapter_response = update_chapter(
            id=chapter["id"],
            ts_id=ctx.user.id,
            ts_name=ctx.user.name,
            ts_at=int(arrow.utcnow().timestamp()),
            should_qced=1 if pred == "Yes" else 0
        )
        if update_chapter_response.status_code != 200:
            await ctx.respond(f"An error occured. Code: {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
            return

        await ctx.respond(f"`Chapter {chap_num}`: `{round(chapter['tl_bytes'] / 1000)}` kilobyte.\n"
                          f"Fee: {calculate_pay(chapter['tl_bytes'], 1 if pred == 'Yes' else 0)} $")
        await message_related_channel(self.bot, role, chap_num, "TS", ctx, pred, round(chapter['tl_bytes'] / 1000))
    
    @discord.slash_command(description="Delete the typeset that you have entered to the bot.")
    async def deletets(
        self, 
        ctx: discord.ApplicationContext, 
        role: discord.Option(Role, name="serie_role", description="Serie role.", required=True),
        chap_num: discord.Option(float, name="chapter_number", description="Chapter number. (like '7' or 7.5')", required=True)
    ):
        serie, chapter = find_serie_and_chapter(role.id, chap_num)
        if not serie:
            await ctx.respond("There is no such serie.")
            return
        if not chapter:
            await ctx.respond("There is no typeset found!")
            return

        update_chapter_response = update_chapter(
            id=chapter["id"],
            ts_id=0,
            ts_name="None",
            ts_at=0,
            should_qced=0
        )
        if update_chapter_response.status_code != 200:
            await ctx.respond(f"An error occured. Code: {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
            return

        await ctx.respond(f"`{role.name}` `{chap_num}` is deleted.")
    
    @discord.slash_command(name="proofread", description="Calculate the proofreading fee and enter it into the bot's database.")
    async def enterpr(
        self, 
        ctx: discord.ApplicationContext, 
        role: discord.Option(Role, name="serie_role", description="Serie role.", required=True),
        chap_num: discord.Option(float, name="chapter_number", description="Chapter number. (like '7' or 7.5')", required=True)
    ):
        serie, chapter = find_serie_and_chapter(role.id, chap_num)
        if not chapter:
            await ctx.respond("There is no such serie or chapter!", ephemeral=True)
            return

        time_str = datetime.now().strftime(r"%d/%m/%Y - %H.%M")

        update_serie_response = update_serie(
            id=serie["id"],
            pred=chap_num,
            time_pr=time_str
        )

        if update_serie_response.status_code != 200:
            await ctx.respond(f"An error occured. Code:  {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
            return

        update_chapter_response = update_chapter(
            id=chapter["id"],
            pr_id=ctx.user.id,
            pr_name=ctx.user.name,
            pr_at=int(arrow.utcnow().timestamp())
        )
        if update_chapter_response.status_code != 200:
            await ctx.respond(f"An error occured. Code:  {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
            return

        await ctx.respond(f"`{role.name}` `{chap_num}` is entered as {PR_FEE}$ to database.")
        await message_related_channel(self.bot, role, chap_num, "PR", ctx)
        
    @discord.slash_command(description="Delete a pr.")
    async def deletepr(
        self, 
        ctx: discord.ApplicationContext, 
        role: discord.Option(Role, name="serie_role", description="Serie role.", required=True), 
        chap_num: discord.Option(float, name="chapter_number", description="Chapter number. (like '7' or 7.5')", required=True)
    ):
        serie, chapter = find_serie_and_chapter(role.id, chap_num)
        if not serie:
            await ctx.respond("There is no such serie.")
            return
        if not chapter:
            await ctx.respond("There is no proofread found!")
            return

        update_chapter_response = update_chapter(
            id=chapter["id"],
            pr_id=0,
            pr_name="",
            pr_at=0
        )
        if update_chapter_response.status_code != 200:
            await ctx.respond(f"An error occured while updating the serie. Code: {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
            return
        
        await ctx.respond(f"`{role.name}` `{chap_num}` is deleted.")
        
    @discord.slash_command(name="clean_fees", description="See all clean fees.")
    @discord.default_permissions(administrator=True)
    async def cln_fees(self, ctx: discord.ApplicationContext):
        await ctx.defer()
        
        with open("mp/clean_fees.json", "r") as file:
            fees = loads(file.read())

        texts = []

        for ID, fee in fees.items():
            texts.append(f"<@&{ID}>: {fee}$")
            if len(texts) == 15:
                await ctx.send_followup("\n".join(texts))
                texts.clear()
                
        await ctx.send_followup("\n".join(texts))
        
    @discord.slash_command(name="delete_clean_fee", description="Delete a clean fee.")
    @discord.default_permissions(administrator=True)
    async def del_clean_fee(self, ctx: discord.ApplicationContext, role: discord.Option(Role, name="serie_role", description="Serie role.", required=True)):
        with open("mp/clean_fees.json", "r") as file:
            fees = loads(file.read())

        fee = fees.get(str(role.id))
        if fee is None:
            await ctx.respond("There is no clean fee for this serie.")
            return
        
        del fees[str(role.id)]
        
        with open("mp/clean_fees.json", "w") as file:
            file.write(dumps(fees))
            
        await ctx.respond(f"`{role.name}` clean fee is deleted.")
    
    @discord.slash_command(name="add_clean_fee", description="Add a clean fee to a serie.")
    @discord.default_permissions(administrator=True)
    async def add_cln_fee(
        self, 
        ctx: discord.ApplicationContext, 
        role: discord.Option(Role, name="serie_role", description="Serie role.", required=True),
        fee: discord.Option(float, name="fee", description="Serie clean fee. (like '7' or 7.5')", required=True)
    ):
        with open("mp/clean_fees.json") as file:
            fees = loads(file.read())

        fees[str(role.id)] = fee

        with open("mp/clean_fees.json", "w") as file:
            file.write(dumps(fees))

        await ctx.respond(f"`{role.name}` fee is `{fee}$` now.")
    
    @discord.slash_command(name="clean", description="Calculate the cleaning fee and enter it into the bot's database.")
    async def calc_clean(
        self, 
        ctx: discord.ApplicationContext,
        role: discord.Option(Role, name="serie_role", description="Serie role.", required=True),
        chap_num: discord.Option(float, name="chapter_number", description="Chapter number. (like '7' or 7.5')", required=True)
    ):
        with open("mp/clean_fees.json") as file:
            fees = loads(file.read())

        if fees.get(str(role.id)) is None:
            await ctx.respond("There is no clean fee for this serie. Please ask to an admin.")
            return

        serie, chapter = find_serie_and_chapter(role.id, chap_num)
        if not chapter:
            await ctx.respond("There is no such serie or chapter!", ephemeral=True)
            return

        time_str = datetime.now().strftime(r"%d/%m/%Y - %H.%M")

        update_serie_response = update_serie(
            id=serie["id"],
            cleaned=chap_num,
            time_cln=time_str
        )
        if update_serie_response.status_code != 200:
            await ctx.respond(f"An error occured while updating the serie. Code: {update_serie_response.status_code}\n```\n{update_serie_response.text}```")
            return

        update_chapter_response = update_chapter(
            id=chapter["id"],
            clnr_id=ctx.user.id,
            clnr_name=ctx.user.name,
            clnr_at=int(arrow.utcnow().timestamp())
        )
        if update_chapter_response.status_code != 200:
            await ctx.respond(f"An error occured while updating the serie. Code: {update_chapter_response.status_code}\n```\n{update_chapter_response.text}```")
            return

        await ctx.respond(f"`{role.name}` `{chap_num}` is added to database. Fee: `{fees[str(role.id)]}$`")
        await message_related_channel(self.bot, role, chap_num, "CLN", ctx)
        
    @discord.slash_command(description="Delete a clean.")
    async def delete_clean(
        self, 
        ctx: discord.ApplicationContext,
        role: discord.Option(Role, name="serie_role", description="Serie role.", required=True),
        chap_num: discord.Option(float, name="chapter_number", description="Chapter number. (like '7' or 7.5')", required=True)
    ):
        serie, chapter = find_serie_and_chapter(role.id, chap_num)
        if not serie:
            await ctx.respond("There is no such serie.")
            return
        if not chapter:
            await ctx.respond("There is no clean found!")
            return
        update_chapter_response = update_chapter(
            id=chapter["id"],
            clnr_id=0,
            clnr_name="",
            clnr_at=0
        )
        if update_chapter_response.status_code != 200:
            await ctx.respond(f"An error occured while updating the serie. Code: {update_chapter_response.status_code}\n```\n{update_chapter_response.text}```")
            return

        await ctx.respond(f"`{role.name}` `{chap_num}` deleted.")
        
    @discord.slash_command(name="qc", description="Calculate the QC fee and enter it into the bot's database.")
    async def calc_qc(
        self, 
        ctx: discord.ApplicationContext,
        role: discord.Option(Role, name="serie_role", description="Serie role.", required=True),
        chap_num: discord.Option(float, name="chapter_number", description="Chapter number. (like '7' or 7.5')", required=True)
    ):
        serie, chapter = find_serie_and_chapter(role.id, chap_num)
        if not chapter:
            await ctx.respond("There is no such serie or chapter!", ephemeral=True)
            return

        time_str = datetime.now().strftime(r"%d/%m/%Y - %H.%M")

        update_serie_response = update_serie(
            id=serie["id"],
            last_qced=chap_num,
            time_qc=time_str
        )
        if update_serie_response.status_code != 200:
            await ctx.respond(f"An error occured while updating the serie. Code: {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
            return

        update_chapter_response = update_chapter(
            id=chapter["id"],
            qc_id=ctx.user.id,
            qc_name=ctx.user.name,
            qc_at=int(arrow.utcnow().timestamp())
        )
        if update_chapter_response.status_code != 200:
            await ctx.respond(f"An error occured while updating the serie. Code: {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
            return

        await ctx.respond(f"`{role.name}` `{chap_num}` added `{QC_FEE}$`.")
        await message_related_channel(self.bot, role, chap_num, "QC", ctx)
    
    @discord.slash_command(description="Delete a qc.")
    async def sonkontrolsil(
        self, 
        ctx: discord.ApplicationContext,
        role: discord.Option(Role, name="serie_role", description="Serie role.", required=True),
        chap_num: discord.Option(float, name="chapter_number", description="Chapter number. (like '7' or 7.5')", required=True)

    ):
        serie, chapter = find_serie_and_chapter(role.id, chap_num)
        if not serie:
            await ctx.respond("There is no such serie.")
            return
        if not chapter:
            await ctx.respond("There is no QC found!")
            return
        update_chapter_response = update_chapter(
            id=chapter["id"],
            qc_id=0,
            qc_name="",
            qc_at=0
        )
        if update_chapter_response.status_code != 200:
            await ctx.respond(f"An error occured while updating the serie. Code: {serie_update_resp.status_code}\n```\n{serie_update_resp.text}```")
            return

        await ctx.respond(f"`{role.name}` `{chap_num}` deleted.")
    
    # noinspection PyMethodMayBeStatic
    def strip_data(self, txt: str):
        data = txt

        data = data.replace("***", "")
        data = data.replace("**", "")
        data = data.replace("//", "")
        data = data.replace("\n", "")
        data = data.replace("\r", "")
        data = data.replace("ST: ", "")
        data = data.replace("OT: ", "")

        spaces = set()
        counter = 0
        for char in data:
            if char == " ":
                counter += 1
            else:
                spaces.add(counter)
                counter = 0

        for space in spaces:
            if space >= 2:
                data = data.replace(" " * space, "")

        return data

    # noinspection PyChainedComparisons
    @staticmethod
    def check_if_work_is_in_between_tstamps(chapter, field: str, lb, ub) -> bool:
        if chapter[field] >= lb and chapter[field] <= ub:
            return True
        
        return False

    @discord.slash_command(description="How rich are you?")
    async def richness(
        self, 
        ctx: discord.ApplicationContext, 
        member: Member = None
    ):
        membr = member
        if not membr:
            membr = ctx.user
            
        now_utc_ts = int(arrow.utcnow().timestamp())

        with rget(f"{API}/series/get/") as response:
            series = loads(response.text)["series"]

        with rget(f"{API}/payperiods/get/last-period") as response:
            if response.status_code != 200:
                await ctx.respond(f"An error occured. Code: {response.status_code}\n```{response.text}```")
                return

            period = loads(response.text)["pay_period"]

        with rget(f"{API}/chapters/get/between-dates/?lower_bound={int(period['created_at'])}&upper_bound={now_utc_ts}") as response:
            results = loads(response.text).get("results")
            if not results:
                await ctx.respond("There is no chapters for this pay period!")
                return

        total_fee = 0

        tl_fields: list[EmbedField] = []
        pr_fields: list[EmbedField] = []
        cln_fields: list[EmbedField] = []
        ts_fields: list[EmbedField] = []
        qc_fields: list[EmbedField] = []

        with open("mp/clean_fees.json") as file:
            clean_fees = loads(file.read())

        ctrl1 = "No control needed."
        ctrl2 = "Control needed."

        for chapter in results:
            if chapter["tl_id"] == membr.id and chapter["should_pred"] == 0 and self.check_if_work_is_in_between_tstamps(chapter, "tl_at", period['created_at'], now_utc_ts):
                tl_fields.append(
                    EmbedField(
                        name=f"{chapter['serie_name']} {chapter['chapter_num']}",
                        value=f"{ctrl1} Size: `{round(chapter['tl_bytes'] / 1000)} KB` Fee: `{calculate_pay(chapter['tl_bytes'], 0)} $`",
                        inline=False
                    )
                )
                total_fee += calculate_pay(chapter['tl_bytes'], 0)
            if chapter["tl_id"] == membr.id and chapter["should_pred"] == 1 and self.check_if_work_is_in_between_tstamps(chapter, "tl_at", period['created_at'], now_utc_ts):
                tl_fields.append(
                    EmbedField(
                        name=f"{chapter['serie_name']} {chapter['chapter_num']}",
                        value=f"{ctrl2} Size: `{round(chapter['tl_bytes'] / 1000)} KB` Fee: `{calculate_pay(chapter['tl_bytes'], 1)}`",
                        inline=False
                    )
                )
                total_fee += calculate_pay(chapter['tl_bytes'], 1)

            if chapter["pr_id"] == membr.id and self.check_if_work_is_in_between_tstamps(chapter, "pr_at", period['created_at'], now_utc_ts):
                pr_fields.append(
                    EmbedField(
                        name=f"{chapter['serie_name']} {chapter['chapter_num']}",
                        value=f"Fee: `{PR_FEE} $`",
                        inline=False
                    )
                )
                total_fee += PR_FEE

            if chapter["clnr_id"] == membr.id and self.check_if_work_is_in_between_tstamps(chapter, "clnr_at", period['created_at'], now_utc_ts):
                fee = None
                for s in series:
                    if s["id"] == chapter["serie_id"]:
                        fee = clean_fees[str(s["role_id"])]
                cln_fields.append(
                    EmbedField(
                        name=f"{chapter['serie_name']} {chapter['chapter_num']}",
                        value=f"Fee: `{fee} $`",
                        inline=False
                    )
                )
                total_fee += fee

            if (chapter["ts_id"] == membr.id or chapter["ts_name"] == membr.name) and chapter["should_qced"] == 0  and self.check_if_work_is_in_between_tstamps(chapter, "ts_at", period['created_at'], now_utc_ts):
                ts_fields.append(
                    EmbedField(
                        name=f"{chapter['serie_name']} {chapter['chapter_num']}",
                        value=f"{ctrl1} Size: `{round(chapter['tl_bytes'] / 1000)} KB` Fee: `{calculate_pay(chapter['tl_bytes'], 0)} $`",
                        inline=False
                    )
                )
                total_fee += calculate_pay(chapter['tl_bytes'], 0)
            if (chapter["ts_id"] == membr.id or chapter["ts_name"] == membr.name) and chapter["should_qced"] == 1 and self.check_if_work_is_in_between_tstamps(chapter, "ts_at", period['created_at'], now_utc_ts):
                ts_fields.append(
                    EmbedField(
                        name=f"{chapter['serie_name']} {chapter['chapter_num']}",
                        value=f"{ctrl2} Size: `{round(chapter['tl_bytes'] / 1000)} KB` Fee: `{calculate_pay(chapter['tl_bytes'], 1)}`",
                        inline=False
                    )
                )
                total_fee += calculate_pay(chapter['tl_bytes'], 1)

            if chapter["qc_id"] == membr.id and self.check_if_work_is_in_between_tstamps(chapter, "qc_at", period['created_at'], now_utc_ts):
                qc_fields.append(
                    EmbedField(
                        name=f"{chapter['serie_name']} {chapter['chapter_num']}",
                        value=f"Fee: `{QC_FEE} $`",
                        inline=False
                    )
                )
                total_fee += QC_FEE

        tl_embeds = generate_embeds_from_fields(tl_fields, membr, "Translations")
        pr_embeds = generate_embeds_from_fields(pr_fields, membr, "Proofreads")
        cln_embeds = generate_embeds_from_fields(cln_fields, membr, "Cleans")
        ts_embeds = generate_embeds_from_fields(ts_fields, membr, "Typesets")
        qc_embeds = generate_embeds_from_fields(qc_fields, membr, "Quality Checks")

        initial_embed = Embed(title=f"{membr.name}")
        total_work_count = sum([len(tl_fields), len(pr_fields), len(cln_fields), len(ts_fields), len(qc_embeds)])
        initial_embed.add_field(
            name=f"Total worked chapters:",
            value=f"`{total_work_count}`"
        )
        initial_embed.add_field(
            name="Total fee:",
            value=f"`{round(total_fee, 2)} $`"
        )
        if total_work_count:
            initial_embed.set_footer(text="Well done!", icon_url=DOLLAR_URL)  # Insert dollar image here
        else:
            initial_embed.set_footer(text="You poor thing!", icon_url=POOP_URL)  # Inser poop image here

        tl_embeds_2 = split_a_page(tl_embeds)
        tl_embeds_3 = split_a_page(tl_embeds_2)
        pr_embeds_2 = split_a_page(pr_embeds)
        pr_embeds_3 = split_a_page(pr_embeds_2)
        cln_embeds_2 = split_a_page(cln_embeds)
        cln_embeds_3 = split_a_page(cln_embeds_2)
        ts_embeds_2 = split_a_page(ts_embeds)
        ts_embeds_3 = split_a_page(ts_embeds_2)
        qc_embeds_2 = split_a_page(qc_embeds)
        qc_embeds_3 = split_a_page(qc_embeds_2)

        pages = [Page(embeds=[initial_embed])]

        if tl_embeds:
            pages.append(Page(embeds=tl_embeds))
        if tl_embeds_2:
            pages.append(Page(embeds=tl_embeds_2))
        if tl_embeds_3:
            pages.append(Page(embeds=tl_embeds_3))
        if pr_embeds:
            pages.append(Page(embeds=pr_embeds))
        if pr_embeds_2:
            pages.append(Page(embeds=pr_embeds_2))
        if pr_embeds_3:
            pages.append(Page(embeds=pr_embeds_3))
        if cln_embeds:
            pages.append(Page(embeds=cln_embeds))
        if cln_embeds_2:
            pages.append(Page(embeds=cln_embeds_2))
        if cln_embeds_3:
            pages.append(Page(embeds=cln_embeds_3))
        if ts_embeds:
            pages.append(Page(embeds=ts_embeds))
        if ts_embeds_2:
            pages.append(Page(embeds=ts_embeds_2))
        if ts_embeds_3:
            pages.append(Page(embeds=ts_embeds_3))
        if qc_embeds:
            pages.append(Page(embeds=qc_embeds))
        if qc_embeds_2:
            pages.append(Page(embeds=qc_embeds_2))
        if qc_embeds_3:
            pages.append(Page(embeds=qc_embeds_3))

        paginator = Paginator(
            pages,
            show_disabled=False,
            author_check=True,
            disable_on_timeout=True,
            timeout=500
        )

        await paginator.respond(ctx.interaction, ephemeral=True)
         
    @discord.slash_command(name="print", description="Print all period fees to an excel file.")
    @discord.default_permissions(administrator=True)
    async def print(
            self,
            ctx: discord.ApplicationContext,
            period_id: discord.Option(
                int,
                name="period_id",
                description="ID of the requested payment period. If not entered, the last period is counted.",
                required=False
            )
    ):
        await ctx.defer()

        period = None
        if not period_id:
            with rget(f"{API}/payperiods/get/last-period") as response:
                if response.status_code != 200:
                    await ctx.respond(f"An error occured while searching periods. Code:{response.status_code}\n```{response.text}```")
                    return

                period = loads(response.text)["pay_period"]
        else:
            with rget(f"{API}/payperiods/get/{period_id}") as response:
                if response.status_code != 200:
                    await ctx.respond(
                        f"An error occured while searching periods. Code:{response.status_code}\n```{response.text}```")
                    return

                period = loads(response.text)["pay_period"]

        file_name = f"database - period {period['id']}"

        with rget(f"{API}/series/get/") as response:
            series = loads(response.text)["series"]
            
        upper_bound = int(period['closed_at']) if period['closed_at'] else int(arrow.utcnow().timestamp())

        with rget(f"{API}/chapters/get/between-dates/?lower_bound={int(period['created_at'])}&upper_bound={upper_bound}") as response:
            results = loads(response.text)["results"]
            if not results:
                await ctx.respond("There is no chapters for this payment period!")
                return

        with open("mp/clean_fees.json") as file:
            fees = loads(file.read())

        chapters_data: list[ChapterAndFees] = []
        for ch in results:
            uwu = ChapterAndFees(
                chapter=ch,
                tl_fee=calculate_pay(ch["tl_bytes"], ch["should_pred"]) if ch["tl_id"] and self.check_if_work_is_in_between_tstamps(ch, "tl_at", period['created_at'], upper_bound) else None,
                pr_fee=PR_FEE if ch["pr_id"] and self.check_if_work_is_in_between_tstamps(ch, "pr_at", period['created_at'], upper_bound) else None,
                cln_fee=fees.get(str(find_serie_by_id(series, ch["serie_id"])["role_id"])) if ch["clnr_id"] and self.check_if_work_is_in_between_tstamps(ch, "clnr_at", period['created_at'], upper_bound) else None,
                ts_fee=calculate_pay(ch["tl_bytes"], ch["should_qced"]) if ch["ts_id"] and self.check_if_work_is_in_between_tstamps(ch, "ts_at", period['created_at'], upper_bound) else None,
                qc_fee=QC_FEE if ch["qc_id"] and self.check_if_work_is_in_between_tstamps(ch, "qc_at", period['created_at'], upper_bound) else None
            )

            if uwu.check_if_any_fee_exist:
                chapters_data.append(uwu)

        personnel_ids = extract_all_personnel_ids_chapters(chapters_data)

        workbook = Workbook(f'mp\\{file_name}.xlsx')
        personnel_names_and_sheets = match_personnel_ids_with_names_and_add_worksheet(personnel_ids, self.bot, workbook)
        merged_format = workbook.add_format({"align": "center"})

        for key, val in personnel_names_and_sheets.items():
            prepare_sheet(val[1], merged_format)
            user_works = split_all_work(chapters_data, key, series, fees, period["created_at"], upper_bound)
            for i in range(len(user_works[0])):
                val[1].write(3 + i, 0, user_works[0][i][0])
                val[1].write(3 + i, 1, round(user_works[0][i][1]/1000))
                val[1].write_formula(3 + i, 2, "=IFS(B4=4, 0.45, B4=2, 0.35, B4=3, 0.4, B4=5, 0.5, B4=6, 0.55, B4=7, 0.6, B4=8, 0.65, B4=9, 0.7, B4=10, 0.75, B4>=11, 0.8, B4<=1, 0.35)")
            for i in range(len(user_works[1])):
                val[1].write(3 + i, 3, user_works[1][i][0])
                val[1].write(3 + i, 4, round(user_works[1][i][1] / 1000))
                val[1].write_formula(3 + i, 5, "=IFS(E4=4, 0.3, E4=2, 0.2, E4=3, 0.25, E4=5, 0.35, E4=6, 0.4, E4=7, 0.45, E4=8, 0.5, E4=9, 0.55, E4=10, 0.6, E4>=11, 0.65, E4<=1, 0.2)")
            for i in range(len(user_works[2])):
                val[1].write(3 + i, 6, user_works[2][i][0])
                val[1].write(3 + i, 7, user_works[2][i][1])
            for i in range(len(user_works[3])):
                pass

        workbook.close()
        await ctx.respond("**Success!**", file=File(f'mp\\{file_name}.xlsx'))
        os.remove(f'mp\\{file_name}.xlsx')


def setup(bot: discord.Bot):
    bot.add_cog(MaasHesap(bot))
